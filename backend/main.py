import email
import hashlib
import hmac
import secrets
import time
import threading
from concurrent.futures import ThreadPoolExecutor
from backend.generator import generate
from collections import defaultdict, deque
import imaplib
import os
import re
import smtplib
import ssl
from datetime import datetime, timezone
from email.header import decode_header, make_header
from email.message import EmailMessage
from email.utils import getaddresses, make_msgid, parsedate_to_datetime
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Request, Response
from fastapi.responses import JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import Boolean, DateTime, ForeignKey, Integer, String, Text, create_engine, func, or_, select
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker

ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / '.env')
engine = create_engine(os.getenv('DATABASE_URL', 'postgresql+psycopg://localhost/cerato_crm'))
Session = sessionmaker(engine, expire_on_commit=False)
def now(): return datetime.now(timezone.utc)
class Base(DeclarativeBase): pass
class Lead(Base):
    __tablename__ = 'leads'
    id: Mapped[int] = mapped_column(primary_key=True)
    company: Mapped[str] = mapped_column(String, index=True)
    city: Mapped[str | None]; category: Mapped[str | None]; service_types: Mapped[str | None]
    website: Mapped[str | None]; employees: Mapped[str | None]; linkedin: Mapped[str | None]
    contact_name: Mapped[str | None]; person_email: Mapped[str | None]; company_email: Mapped[str | None]
    status: Mapped[str] = mapped_column(String, default='new', index=True)
    notes: Mapped[str | None] = mapped_column(Text); research: Mapped[str | None] = mapped_column(Text)
    research_sources: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
class Message(Base):
    __tablename__ = 'messages'
    id: Mapped[int] = mapped_column(primary_key=True)
    lead_id: Mapped[int] = mapped_column(ForeignKey('leads.id'), index=True)
    direction: Mapped[str] = mapped_column(String, default='outbound')
    status: Mapped[str] = mapped_column(String, default='draft')
    sender: Mapped[str] = mapped_column(String)
    recipient: Mapped[str] = mapped_column(String)
    subject: Mapped[str] = mapped_column(String)
    body: Mapped[str] = mapped_column(Text)
    message_id: Mapped[str | None] = mapped_column(String, unique=True)
    in_reply_to: Mapped[str | None] = mapped_column(String)
    error: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now)
    sent_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True))
class Suppression(Base):
    __tablename__ = 'suppressions'
    email: Mapped[str] = mapped_column(String, primary_key=True)
    reason: Mapped[str | None] = mapped_column(Text)
class EmailTemplate(Base):
    __tablename__ = 'email_templates'
    id: Mapped[int] = mapped_column(primary_key=True)
    subject: Mapped[str] = mapped_column(String)
    body: Mapped[str] = mapped_column(Text)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now, onupdate=now)

DEFAULT_TEMPLATE = """Hi {{first_name}},

{{personalized_sentence}}

I’m Stefan from Cerato Systems. We help logistics and transport companies reduce the manual work between their TMS/ERP, email, documents, and spreadsheets.

I’d love to record a short video for {{company_name}} showing 3 places where that kind of manual work could be automated, using the systems you already have. It’s completely free, and yours to keep either way.

Want me to record the walkthrough?

Best,
Stefan
CTO, Cerato Systems"""
Base.metadata.create_all(engine)
app = FastAPI(title='Cerato CRM')
app.add_middleware(CORSMiddleware, allow_origins=['http://localhost:5173','http://127.0.0.1:5173'], allow_methods=['*'], allow_headers=['*'])

# Public deployments require a configured password hash and session key.
AUTH_HASH = os.getenv('CRM_PASSWORD_HASH', '')
SESSION_KEY = os.getenv('CRM_SESSION_KEY', '')
SECURE_COOKIES = os.getenv('CRM_SECURE_COOKIES', 'false').lower() == 'true'
if bool(AUTH_HASH) != bool(SESSION_KEY):
    raise RuntimeError('Set both CRM_PASSWORD_HASH and CRM_SESSION_KEY')
LOGIN_ATTEMPTS = defaultdict(deque)
SESSION_MAX_AGE = 12 * 60 * 60

def session_valid(token):
    try:
        timestamp, signature = token.split('.', 1)
        age = time.time() - int(timestamp)
        expected = hmac.new(SESSION_KEY.encode(), timestamp.encode(), hashlib.sha256).hexdigest()
        return 0 <= age <= SESSION_MAX_AGE and hmac.compare_digest(signature, expected)
    except (ValueError, AttributeError):
        return False

def password_valid(password):
    try:
        algorithm, iterations, salt, expected = AUTH_HASH.split('$')
        if algorithm != 'pbkdf2_sha256': return False
        actual = hashlib.pbkdf2_hmac('sha256', password.encode(), bytes.fromhex(salt), int(iterations)).hex()
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError): return False

@app.middleware('http')
async def require_auth(request: Request, call_next):
    if not AUTH_HASH: return await call_next(request)
    path = request.url.path
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        origin = request.headers.get('origin')
        expected_origin = os.getenv('CRM_ORIGIN', '')
        if origin and origin != expected_origin:
            return JSONResponse({'detail':'Origin not allowed'}, status_code=403)
        if request.headers.get('sec-fetch-site') == 'cross-site':
            return JSONResponse({'detail':'Cross-site request rejected'}, status_code=403)
    if path not in ('/api/auth/status', '/api/auth/login') and not session_valid(request.cookies.get('cerato_session', '')):
        return JSONResponse({'detail':'Login required'}, status_code=401)
    return await call_next(request)

class LoginInput(BaseModel):
    password: str

@app.get('/api/auth/status')
def auth_status(request: Request):
    return {'required': bool(AUTH_HASH), 'authenticated': not AUTH_HASH or session_valid(request.cookies.get('cerato_session', ''))}

@app.post('/api/auth/login')
def login(data: LoginInput, request: Request, response: Response):
    if not AUTH_HASH: return {'ok': True}
    # Nginx should overwrite X-Real-IP rather than forwarding an untrusted header.
    ip = request.headers.get('x-real-ip', request.client.host if request.client else 'unknown')
    attempts = LOGIN_ATTEMPTS[ip]; current = time.monotonic()
    while attempts and attempts[0] < current - 900: attempts.popleft()
    if len(attempts) >= 10: raise HTTPException(429, 'Too many attempts. Try again later.')
    attempts.append(current)
    if not password_valid(data.password): raise HTTPException(401, 'Incorrect password')
    attempts.clear()
    timestamp = str(int(time.time()))
    signature = hmac.new(SESSION_KEY.encode(), timestamp.encode(), hashlib.sha256).hexdigest()
    response.set_cookie('cerato_session', timestamp+'.'+signature, max_age=SESSION_MAX_AGE, httponly=True, secure=SECURE_COOKIES, samesite='strict', path='/')
    return {'ok':True}

@app.post('/api/auth/logout')
def logout(response: Response):
    response.delete_cookie('cerato_session', path='/', secure=SECURE_COOKIES, httponly=True, samesite='strict')
    return {'ok':True}

def serialize(obj):
    return {c.name: getattr(obj, c.name) for c in obj.__table__.columns}
def mailbox(key):
    if key not in ('stefan', 'adrija'): raise HTTPException(400, 'Choose Stefan or Adrija')
    address, password = os.getenv(key.upper()+'_EMAIL','').strip(), os.getenv(key.upper()+'_PASSWORD','')
    if not address or not password: raise HTTPException(400, f'{key.title()} mailbox is not configured in .env')
    return address, password
def valid_address(value): return bool(re.fullmatch(r'[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+', value or ''))
def get_lead(db, id):
    obj = db.get(Lead,id)
    if not obj: raise HTTPException(404, 'Lead not found')
    return obj
class LeadUpdate(BaseModel):
    status: str | None = None
    notes: str | None = None
    research: str | None = None
    research_sources: str | None = None
class DraftInput(BaseModel):
    lead_id: int
    sender: str
    recipient: str
    subject: str
    body: str
class SuppressInput(BaseModel):
    email: str
    reason: str = 'Opted out'
class TemplateInput(BaseModel):
    subject: str = ''
    body: str

def template_id(sender: str):
    if sender not in ('stefan', 'adrija'): raise HTTPException(400, 'Choose Stefan or Andrija')
    return 1 if sender == 'stefan' else 2

def default_template_for(sender: str):
    if sender == 'stefan': return DEFAULT_TEMPLATE
    return DEFAULT_TEMPLATE.replace('I’m Stefan from Cerato Systems.', 'I’m Andrija from Cerato Systems.').replace('Best,\nStefan\nCTO, Cerato Systems', 'Best,\nAndrija\nCerato Systems')

@app.get('/api/template')
def get_template(sender: str = 'stefan'):
    key = template_id(sender)
    with Session.begin() as db:
        template = db.get(EmailTemplate, key)
        if not template:
            template = EmailTemplate(id=key, subject='', body=default_template_for(sender))
            db.add(template)
            db.flush()
        return {**serialize(template), 'sender': sender}

@app.put('/api/template')
def update_template(data: TemplateInput, sender: str = 'stefan'):
    key = template_id(sender)
    if not data.body.strip(): raise HTTPException(400, 'Template body cannot be empty')
    if '\n' in data.subject or '\r' in data.subject: raise HTTPException(400, 'Subject must be a single line')
    with Session.begin() as db:
        template = db.get(EmailTemplate, key)
        if not template:
            template = EmailTemplate(id=key, subject=data.subject, body=data.body)
            db.add(template)
        else:
            template.subject = data.subject
            template.body = data.body
        db.flush()
        return {**serialize(template), 'sender': sender}

@app.get('/api/health')
def health(): return {'ok': True}
@app.get('/api/settings')
def settings(): return {'mailboxes': [{'id': x, 'email': os.getenv(x.upper()+'_EMAIL',''), 'configured': bool(os.getenv(x.upper()+'_EMAIL') and os.getenv(x.upper()+'_PASSWORD'))} for x in ('stefan','adrija')]}
@app.post('/api/import')
def import_leads():
    path = ROOT / 'road_groupage_companies.xlsx'
    if not path.exists(): raise HTTPException(404,'Spreadsheet not found')
    rows = load_workbook(path, read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True)
    added = skipped = 0
    with Session.begin() as db:
        existing = set(db.scalars(select(Lead.company)).all())
        for row in rows:
            if not row[0]: continue
            vals = [str(x).strip() if x is not None else None for x in row[:10]]
            if vals[0] in existing: skipped += 1; continue
            db.add(Lead(**dict(zip(['company','city','category','service_types','website','employees','linkedin','contact_name','person_email','company_email'],vals))))
            existing.add(vals[0]); added += 1
    return {'added':added,'skipped':skipped}
@app.get('/api/leads')
def leads(q: str = '', status: str = '', groupage: bool = False):
    with Session() as db:
        query = select(Lead)
        if q: query = query.where(or_(Lead.company.ilike('%'+q+'%'), Lead.city.ilike('%'+q+'%'), Lead.person_email.ilike('%'+q+'%'), Lead.company_email.ilike('%'+q+'%')))
        if status: query = query.where(Lead.status == status)
        if groupage: query = query.where(Lead.category.ilike('%Groupage by road%'))
        owners = {}
        for lead_id, sender in db.execute(select(Message.lead_id, Message.sender).where(Message.direction == 'outbound').order_by(Message.created_at, Message.id)):
            owners.setdefault(lead_id, sender)
        return [{**serialize(x), 'owner': owners.get(x.id)} for x in db.scalars(query.order_by(Lead.company)).all()]
@app.get('/api/leads/{id}')
def lead(id: int):
    with Session() as db:
        obj=get_lead(db,id)
        history = db.scalars(select(Message).where(Message.lead_id==id).order_by(Message.created_at.desc(), Message.id.desc())).all()
        outbound = [x for x in reversed(history) if x.direction == 'outbound']
        return {**serialize(obj), 'owner': outbound[0].sender if outbound else None, 'messages':[serialize(x) for x in history]}
@app.patch('/api/leads/{id}')
def update_lead(id: int, data: LeadUpdate):
    with Session.begin() as db:
        obj=get_lead(db,id)
        changes=data.model_dump(exclude_unset=True)
        if 'status' in changes and changes['status'] not in ('new','review','simulated','sent','replied','do_not_contact'): raise HTTPException(400,'Invalid status')
        if obj.status == 'do_not_contact' and changes.get('status', 'do_not_contact') != 'do_not_contact': raise HTTPException(409, 'A stopped lead cannot be reactivated here')
        for k,v in changes.items(): setattr(obj,k,v)
        if obj.status == 'do_not_contact':
            for address in (obj.person_email,obj.company_email):
                if valid_address(address) and not db.get(Suppression,address.lower()): db.add(Suppression(email=address.lower(),reason='Lead marked do not contact'))
        return serialize(obj)
@app.get('/api/messages')
def messages(status: str = ''):
    with Session() as db:
        query=select(Message,Lead.company).join(Lead)
        if status: query=query.where(Message.status==status)
        return [{**serialize(m),'company':company} for m,company in db.execute(query.order_by(Message.created_at.desc())).all()]
@app.post('/api/drafts')
def draft(data: DraftInput):
    if data.sender not in ('stefan','adrija') or not valid_address(data.recipient) or not data.subject.strip() or not data.body.strip(): raise HTTPException(400,'Provide a sender, one valid recipient, subject and body')
    with Session.begin() as db:
        lead=db.scalar(select(Lead).where(Lead.id==data.lead_id).with_for_update())
        if not lead: raise HTTPException(404, 'Lead not found')
        if db.scalar(select(Message.id).where(Message.lead_id==lead.id, Message.status.in_(['draft','simulated','sent','sending','uncertain']), Message.direction=='outbound').limit(1)): raise HTTPException(409, 'This lead already has a draft or outgoing message')
        if lead.status in ('do_not_contact','replied'): raise HTTPException(409,'Lead is stopped or has replied')
        if db.get(Suppression,data.recipient.lower()): raise HTTPException(409,'Recipient is suppressed')
        obj=Message(**data.model_dump(),direction='outbound',status='draft'); db.add(obj); db.flush(); lead.status='review'
        return serialize(obj)
@app.put('/api/drafts/{id}')
def update_draft(id: int, data: DraftInput):
    if data.sender not in ('stefan','adrija') or not valid_address(data.recipient) or not data.subject.strip() or not data.body.strip(): raise HTTPException(400,'Provide a sender, one valid recipient, subject and body')
    with Session.begin() as db:
        obj=db.get(Message,id)
        if not obj or obj.status != 'draft' or obj.lead_id != data.lead_id: raise HTTPException(404,'Draft not found')
        lead = get_lead(db, obj.lead_id)
        if data.sender != obj.sender: raise HTTPException(409, 'This draft belongs to a different sender')
        if lead.status in ('do_not_contact','replied') or db.get(Suppression,data.recipient.lower()): raise HTTPException(409,'Lead or recipient is stopped')
        for k,v in data.model_dump().items(): setattr(obj,k,v)
        return serialize(obj)
@app.delete('/api/drafts/{id}')
def delete_draft(id:int):
    with Session.begin() as db:
        obj=db.get(Message,id)
        if not obj or obj.status != 'draft': raise HTTPException(404,'Draft not found')
        lead=get_lead(db,obj.lead_id); db.delete(obj); db.flush()
        if lead.status=='review' and not db.scalar(select(Message.id).where(Message.lead_id==lead.id,Message.status=='draft').limit(1)): lead.status='new'
    return {'ok':True}
@app.post('/api/drafts/{id}/send')
def send(id: int):
    raise HTTPException(403, 'Real email sending is disabled while simulation mode is active')

@app.post('/api/drafts/{id}/simulate')
def simulate(id: int):
    with Session.begin() as db:
        obj = db.scalar(select(Message).where(Message.id == id).with_for_update())
        if not obj or obj.status != 'draft': raise HTTPException(409, 'Message is not an available draft')
        if re.search(r'{{\s*[^{}]+\s*}}', obj.subject + obj.body): raise HTTPException(400, 'Replace all template placeholders before recording a simulation')
        lead = db.scalar(select(Lead).where(Lead.id == obj.lead_id).with_for_update())
        if lead.status in ('replied', 'do_not_contact'): raise HTTPException(409, 'Lead has replied or is stopped')
        if db.get(Suppression, obj.recipient.lower()): raise HTTPException(409, 'Recipient is suppressed')
        if db.scalar(select(Message.id).where(Message.lead_id == lead.id, Message.id != obj.id, Message.direction == 'outbound', Message.status.in_(['simulated', 'sent', 'sending', 'uncertain'])).limit(1)): raise HTTPException(409, 'An outreach message is already recorded for this lead')
        obj.status = 'simulated'
        lead.status = 'simulated'
        return {'ok': True, 'message': serialize(obj)}
class SimulationRevision(BaseModel):
    body: str

@app.patch('/api/simulations/{id}')
def revise_simulation(id: int, data: SimulationRevision):
    if not data.body.strip() or re.search(r'{{\s*[^{}]+\s*}}', data.body):
        raise HTTPException(400, 'Provide a complete message without placeholders')
    with Session.begin() as db:
        obj = db.scalar(select(Message).where(Message.id == id).with_for_update())
        if not obj or obj.status != 'simulated': raise HTTPException(404, 'Simulation not found')
        obj.body = data.body
        return serialize(obj)

# Jobs deliberately stay in memory: a service restart interrupts them instead of resending.
JOBS = {}
JOB_LOCK = threading.Lock()
JOB_POOL = ThreadPoolExecutor(max_workers=1)
class GenerateInput(BaseModel):
    sender: str

def run_generation(id: int, sender: str, lead_snapshot: dict, template: dict):
    try:
        result = generate(lead_snapshot, template, sender)
        sources = result.get('sources', [])
        if not isinstance(sources, list): raise ValueError('Invalid research sources')
        urls = [u.strip() for u in sources if isinstance(u, str) and re.match(r'^https?://[^\s]+$', u.strip())]
        research = str(result.get('research', '')).strip()[:12000]
        if not research or not urls: raise ValueError('No sourced research returned')
        # Research is retained even when the agent could not verify a recipient.
        with Session.begin() as db:
            obj = db.scalar(select(Lead).where(Lead.id == id).with_for_update())
            if not obj or obj.status != 'new' or db.scalar(select(Message.id).where(Message.lead_id == id).limit(1)):
                raise ValueError('Lead changed while research was running; no draft was created')
            obj.research = research
            obj.research_sources = '\n'.join(urls[:20])
            if result.get('ready') is not True:
                outcome = {'state':'needs_review', 'message':str(result.get('reason') or 'Recipient or research could not be verified')[:1000]}
            else:
                recipient = str(result.get('recipient','')).strip()
                available = [str(x).strip().lower() for x in (obj.person_email, obj.company_email) if x]
                if not valid_address(recipient) or recipient.lower() not in available or db.get(Suppression, recipient.lower()):
                    raise ValueError('The verified recipient does not match an eligible spreadsheet address')
                first = str(result.get('first_name','')).strip()
                sentence = str(result.get('personalized_sentence','')).strip()
                subject = (template['subject'] or str(result.get('subject',''))).strip()
                if not first or not sentence or not subject or '\n' in subject or len(subject)>250: raise ValueError('The generated message is incomplete')
                replacements={'{{first_name}}':first, '{{personalized_sentence}}':sentence, '{{company_name}}':obj.company}
                body=template['body']
                for key,value in replacements.items(): body=body.replace(key,value); subject=subject.replace(key,value)
                if re.search(r'{{\s*[^{}]+\s*}}',body+subject): raise ValueError('The template contains unresolved placeholders')
                message=Message(lead_id=id,direction='outbound',status='draft',sender=sender,recipient=recipient,subject=subject,body=body)
                db.add(message); db.flush(); obj.status='review'
                outcome = {'state':'completed','message':'Draft created. Open the conversation to review it.'}
        with JOB_LOCK: JOBS[id] = outcome
    except Exception as exc:
        with JOB_LOCK: JOBS[id] = {'state':'failed', 'message':str(exc)[:500]}

@app.post('/api/leads/{id}/generate')
def generate_lead(id: int, data: GenerateInput):
    template_id(data.sender)
    with JOB_LOCK:
        if any(x.get('state') == 'running' for x in JOBS.values()): raise HTTPException(409,'Another research job is running; try again when it finishes')
        with Session() as db:
            obj=get_lead(db,id)
            if obj.status != 'new' or db.scalar(select(Message.id).where(Message.lead_id==id).limit(1)): raise HTTPException(409,'Only uncontacted new leads can be researched')
            snapshot={k:getattr(obj,k) for k in ('id','company','city','category','website','linkedin','contact_name','person_email','company_email')}
        template=get_template(data.sender)
        JOBS[id]={'state':'running','message':'Researching this lead and preparing a draft…'}
        JOB_POOL.submit(run_generation,id,data.sender,snapshot,template)
    return JOBS[id]

@app.get('/api/leads/{id}/generation')
def generation_status(id: int):
    with Session() as db: get_lead(db,id)
    with JOB_LOCK: return JOBS.get(id, {'state':'idle','message':''})

@app.post('/api/suppressions')
def suppress(data:SuppressInput):
    address=data.email.strip().lower()
    if not valid_address(address): raise HTTPException(400,'Invalid email')
    with Session.begin() as db:
        if not db.get(Suppression,address): db.add(Suppression(email=address,reason=data.reason))
    return {'ok':True}
def decoded(value):
    try: return str(make_header(decode_header(value or '')))
    except Exception: return value or ''
def plain_body(msg):
    parts=msg.walk() if msg.is_multipart() else [msg]
    for part in parts:
        if part.get_content_type()=='text/plain' and part.get_content_disposition()!='attachment':
            payload=part.get_payload(decode=True)
            if payload: return payload.decode(part.get_content_charset() or 'utf-8',errors='replace')[:50000]
    return ''
@app.post('/api/sync/{sender}')
def sync(sender:str):
    address,password=mailbox(sender); matched=0
    try:
        with imaplib.IMAP4_SSL(os.getenv('IMAP_HOST','imap.purelymail.com'),int(os.getenv('IMAP_PORT','993'))) as imap:
            imap.login(address,password); imap.select('INBOX',readonly=True)
            typ, ids=imap.search(None,'ALL')
            if typ!='OK': raise RuntimeError('Could not search inbox')
            for uid in ids[0].split():
                typ, result=imap.fetch(uid,'(BODY.PEEK[])')
                if typ!='OK' or not result or not isinstance(result[0],tuple): continue
                msg=email.message_from_bytes(result[0][1]); mid=msg.get('Message-ID'); references=(msg.get('In-Reply-To','')+' '+msg.get('References',''))
                if not mid: continue
                with Session.begin() as db:
                    if db.scalar(select(Message.id).where(Message.message_id==mid)): continue
                    outgoing=db.scalars(select(Message).where(Message.sender==sender,Message.direction=='outbound',Message.status.in_(['sent','sending','uncertain']))).all()
                    original=next((m for m in outgoing if m.message_id and m.message_id in references),None)
                    if not original: continue
                    from_addresses=[a.lower() for _,a in getaddresses(msg.get_all('From',[]))]
                    if original.recipient.lower() not in from_addresses: continue
                    lead=get_lead(db,original.lead_id)
                    db.add(Message(lead_id=lead.id,direction='inbound',status='received',sender=sender,recipient=address,subject=decoded(msg.get('Subject')),body=plain_body(msg),message_id=mid,in_reply_to=msg.get('In-Reply-To')))
                    if lead.status!='do_not_contact': lead.status='replied'
                    matched+=1
    except HTTPException: raise
    except Exception as exc: raise HTTPException(502,f'Inbox sync failed: {str(exc)[:300]}')
    return {'matched':matched}
